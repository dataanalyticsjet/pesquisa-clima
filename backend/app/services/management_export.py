from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_STRINGS = {
    "pt": {
        "regional": "Regional",
        "completed": "Formulários concluídos",
        "bases": "Bases cadastradas",
        "base_code": "Código da base",
        "base_name": "Nome oficial da base",
        "generated": "Gerado em (UTC)",
        "note": "Contagens operacionais mostram formulários concluídos por Regional e Base para perfis autorizados da Gestão. A planilha não contém respostas individuais nem dados de identidade.",
    },
    "zh": {
        "regional": "区域",
        "completed": "已完成问卷数",
        "bases": "已注册服务中心数",
        "base_code": "服务中心代码",
        "base_name": "官方名称",
        "generated": "生成时间（UTC）",
        "note": "此表向获授权的管理人员展示各区域和服务中心已完成问卷的运营数量，不包含个人回答或身份信息。",
    },
}


def build_adherence_workbook(payload: dict[str, object], locale: str = "pt") -> bytes:
    """Create an XLSX workbook using only the already-suppressed API payload."""
    labels = _STRINGS[locale if locale in _STRINGS else "pt"]
    regionals = payload["regionals"]
    if not isinstance(regionals, list):
        raise ValueError("The adherence payload has no regional rows")

    workbook = Workbook()
    workbook.properties.creator = "Pesquisa de Engajamento - Voz do Colaborador"
    workbook.properties.title = (
        "Adesão por Regional e Base"
        if locale != "zh"
        else "按区域和服务中心查看参与情况"
    )
    workbook.properties.subject = (
        "Contagens operacionais anônimas de formulários concluídos"
        if locale != "zh"
        else "匿名问卷运营统计"
    )

    regional_sheet = workbook.active
    regional_sheet.title = "区域" if locale == "zh" else "Regional"
    regional_headers = [labels["regional"], labels["completed"], labels["bases"]]
    regional_sheet.append(regional_headers)

    base_sheet = workbook.create_sheet("Base")
    if locale == "zh":
        base_sheet.title = "服务中心"
    base_headers = [
        labels["regional"],
        labels["base_code"],
        labels["base_name"],
        labels["completed"],
    ]
    base_sheet.append(base_headers)

    for regional in regionals:
        regional_name = str(regional["regional_name"])
        regional_count = int(regional["response_count"])
        regional_sheet.append(
            [
                regional_name,
                regional_count,
                int(regional["base_count"]),
            ]
        )
        for base in regional["bases"]:
            base_count = int(base["response_count"])
            base_sheet.append(
                [
                    regional_name,
                    str(base["base_code"]),
                    str(base["base_name"]),
                    base_count,
                ]
            )

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    for sheet in (regional_sheet, base_sheet):
        header_row = 1
        last_table_row = sheet.max_row
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{last_table_row}"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[header_row].height = 30
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="E60012")
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        for row in sheet.iter_rows(min_row=2, max_row=last_table_row):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="FFF5F5")
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        thin_gray = Side(style="thin", color="E5E7EB")
        for row in sheet.iter_rows(min_row=1, max_row=last_table_row):
            for cell in row:
                cell.border = Border(bottom=thin_gray)

        metadata_row = last_table_row + 2
        sheet.cell(metadata_row, 1, labels["generated"])
        sheet.cell(metadata_row, 2, generated_at)
        sheet.cell(
            metadata_row + 1,
            1,
            labels["note"],
        )
        sheet.merge_cells(
            start_row=metadata_row + 1,
            start_column=1,
            end_row=metadata_row + 1,
            end_column=sheet.max_column,
        )
        sheet.cell(metadata_row, 1).font = Font(bold=True, color="666B73")
        sheet.cell(metadata_row, 2).font = Font(color="666B73")
        sheet.cell(metadata_row + 1, 1).font = Font(italic=True, color="666B73", size=10)
        sheet.cell(metadata_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[metadata_row + 1].height = 30

    for column, width in enumerate((24, 28, 22), start=1):
        regional_sheet.column_dimensions[get_column_letter(column)].width = width
    for column, width in enumerate((18, 18, 46, 28), start=1):
        base_sheet.column_dimensions[get_column_letter(column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
