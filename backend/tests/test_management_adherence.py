from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

from app.core import auth as auth_core
from app.core.auth import get_current_user, get_db_session
from app.main import app
from app.repositories import management_repository
from app.services.organization_catalog import REGIONAL_SC_CATALOG


client = TestClient(app)


def _paired_rows_for_regionals(rows):
    base_by_regional = {
        regional.code: regional.service_centers[0].code for regional in REGIONAL_SC_CATALOG
    }
    return [
        (regional_code, base_by_regional.get(regional_code, "UNLISTED_BASE"), count)
        for regional_code, count in rows
    ]


def _base_rows_for_pairs(rows):
    counts = {}
    for _regional_code, base_code, count in rows:
        if base_code is not None:
            counts[base_code] = counts.get(base_code, 0) + count
    return list(counts.items())


class FakeSession:
    def __init__(self, *, total=0, rows=()):
        self.total = total
        self.rows = list(rows)
        self.statements = []

    def scalar(self, statement):
        self.statements.append(statement)
        return self.total

    def execute(self, statement):
        self.statements.append(statement)
        return SimpleNamespace(all=lambda: self.rows)


@pytest.fixture(autouse=True)
def clean_dependencies():
    client.cookies.clear()
    app.dependency_overrides.clear()
    yield
    client.cookies.clear()
    app.dependency_overrides.clear()


def configure_request(
    monkeypatch,
    *,
    roles=("ADMIN",),
    total=0,
    regional_rows=(),
    paired_rows=None,
    survey_exists=True,
):
    user = SimpleNamespace(id=17, status="ACTIVE", access_type="INTERNAL")
    survey = SimpleNamespace(
        id=301,
        code="CLIMATE_2026",
        status="ACTIVE",
        min_group_size=5,
    ) if survey_exists else None
    session = FakeSession(total=total, rows=regional_rows)
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_db_session] = lambda: session
    monkeypatch.setattr(auth_core.identity_repository, "get_user_roles", lambda db, user_id: list(roles))
    monkeypatch.setattr(management_repository, "get_survey_by_code", lambda db, code: survey)
    monkeypatch.setattr(management_repository, "count_anonymous_responses", lambda db, survey_id: total)
    monkeypatch.setattr(
        management_repository,
        "get_anonymous_response_counts_by_regional",
        lambda db, survey_id: list(regional_rows),
    )
    effective_pairs = list(
        _paired_rows_for_regionals(regional_rows) if paired_rows is None else paired_rows
    )
    monkeypatch.setattr(
        management_repository,
        "get_anonymous_response_counts_by_regional_and_base",
        lambda db, survey_id: effective_pairs,
    )
    monkeypatch.setattr(
        management_repository,
        "get_anonymous_response_counts_by_base",
        lambda db, survey_id: _base_rows_for_pairs(effective_pairs),
    )
    return session


def test_regional_count_query_counts_distinct_anonymous_responses_only():
    session = FakeSession(total=0, rows=[("CE", 5), ("GP", 2)])

    result = management_repository.get_anonymous_response_counts_by_regional(session, survey_id=301)

    sql = str(session.statements[0].compile()).lower()
    params = list(session.statements[0].compile().params.values())
    assert result == [("CE", 5), ("GP", 2)]
    assert "anonymous_response_segments" in sql
    assert "anonymous_responses" in sql
    assert "count(distinct anonymous_response_segments.response_id)" in sql
    assert "survey_participation" not in sql
    assert "users" not in sql
    assert "response_answers" not in sql
    assert "response_answer_options" not in sql
    assert "REGIONAL" in params


def test_regional_base_count_query_aggregates_distinct_anonymous_pairs_only():
    session = FakeSession(rows=[("SR", "BNU", 5), ("SR", "BLU", 2)])

    result = management_repository.get_anonymous_response_counts_by_regional_and_base(
        session, survey_id=301
    )

    compiled = session.statements[0].compile()
    sql = str(compiled).lower()
    assert result == [("SR", "BNU", 5), ("SR", "BLU", 2)]
    assert "count(distinct anonymous_responses.response_id)" in sql
    assert "anonymous_response_segments" in sql
    assert "anonymous_responses" in sql
    assert "response_answers" not in sql
    assert "survey_participation" not in sql
    assert "users" not in sql


def test_base_count_query_counts_distinct_anonymous_forms_without_identity_or_answers():
    session = FakeSession(rows=[("ANA", 4), ("BSB", 2)])

    result = management_repository.get_anonymous_response_counts_by_base(session, survey_id=301)

    compiled = session.statements[0].compile()
    sql = str(compiled).lower()
    assert result == [("ANA", 4), ("BSB", 2)]
    assert "count(distinct anonymous_response_segments.response_id)" in sql
    assert "segment_type" in sql
    assert "anonymous_responses" in sql
    assert "survey_participation" not in sql
    assert "users" not in sql
    assert "response_answers" not in sql
    assert "BASE" in compiled.params.values()


def test_total_form_count_uses_distinct_anonymous_response_uuid():
    session = FakeSession(total=11)

    result = management_repository.count_anonymous_responses(session, survey_id=301)

    sql = str(session.statements[0].compile()).lower()
    assert result == 11
    assert "count(distinct anonymous_responses.response_id)" in sql
    assert "survey_participation" not in sql
    assert "users" not in sql


@pytest.mark.parametrize("role", ["ADMIN", "MANAGEMENT", "SURVEY_ADMIN"])
def test_management_roles_can_read_adherence(monkeypatch, role):
    configure_request(monkeypatch, roles=(role,), total=0)

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence")

    assert response.status_code == 200
    assert response.json()["total_regionals"] == 10
    assert response.json()["total_bases"] == 32


@pytest.mark.parametrize("access_type", ["INTERNAL", "EXTERNAL"])
def test_collaborator_only_cannot_read_adherence(monkeypatch, access_type):
    user = SimpleNamespace(id=17, status="ACTIVE", access_type=access_type)
    configure_request(monkeypatch, roles=("COLLABORATOR",))
    app.dependency_overrides[get_current_user] = lambda: user

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence")

    assert response.status_code == 403
    assert response.json() == {"detail": "MANAGEMENT_ACCESS_REQUIRED"}


def test_adherence_requires_session():
    app.dependency_overrides[get_db_session] = FakeSession

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence")

    assert response.status_code == 401


def test_adherence_export_requires_session():
    app.dependency_overrides[get_db_session] = FakeSession

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence/export")

    assert response.status_code == 401


def test_unknown_survey_returns_404(monkeypatch):
    configure_request(monkeypatch, survey_exists=False)

    response = client.get("/api/management/surveys/UNKNOWN/adherence")

    assert response.status_code == 404
    assert response.json() == {"detail": "SURVEY_NOT_FOUND"}


def test_adherence_returns_all_official_regionals_sorted_by_visible_counts(monkeypatch):
    configure_request(monkeypatch, total=13, regional_rows=[("GP", 8), ("CE", 5)])

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()

    assert payload["total_responses"] == 13
    assert payload["total_regionals"] == 10
    assert payload["total_bases"] == 32
    assert payload["regionals_with_results"] == 2
    assert len(payload["regionals"]) == 10
    assert [item["regional_code"] for item in payload["regionals"][:2]] == ["GP", "CE"]
    assert {item["regional_code"] for item in payload["regionals"]} == {
        entry.code for entry in REGIONAL_SC_CATALOG
    }
    assert all(
        item["response_count"] == 0
        for item in payload["regionals"]
        if item["regional_code"] not in {"GP", "CE"}
    )
    assert sum(item["base_count"] for item in payload["regionals"]) == 32
    all_base_codes = [
        base["base_code"] for regional in payload["regionals"] for base in regional["bases"]
    ]
    assert len(all_base_codes) == len(set(all_base_codes)) == 32
    assert all(
        len(item["bases"]) == item["base_count"]
        for item in payload["regionals"]
    )
    by_regional = {item["regional_code"]: item for item in payload["regionals"]}
    assert [base["base_code"] for base in by_regional["PR"]["bases"]] == ["SJS"]
    assert [base["base_code"] for base in by_regional["SR"]["bases"]] == ["BNU", "NSR", "BLU"]
    assert by_regional["GP"]["bases"][0]["response_count"] == 8
    assert by_regional["CE"]["bases"][0]["response_count"] == 5
    gp_bases = {base["base_code"]: base for base in by_regional["GP"]["bases"]}
    assert gp_bases["MAO"]["base_name"] == "AM MAO - AM - FLORES"
    assert gp_bases["RBR"]["base_name"] == "AC RBR - AC SANTA INES"
    forbidden = {"user_id", "email", "response_id", "participation_id", "answer_id"}
    assert forbidden.isdisjoint(payload)
    for regional in payload["regionals"]:
        assert forbidden.isdisjoint(regional)
        assert all(forbidden.isdisjoint(base) for base in regional["bases"])
        assert all("raw_count" not in base for base in regional["bases"])
    assert "adherence_percent" not in payload


def test_operational_regional_counts_are_visible_below_the_analytical_minimum(monkeypatch):
    configure_request(monkeypatch, total=7, regional_rows=[("GP", 5), ("CE", 2)])

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    by_code = {item["regional_code"]: item for item in payload["regionals"]}

    assert payload["total_responses"] == 7
    assert payload["regionals_with_results"] == 1
    assert by_code["GP"]["response_count"] == 5
    assert by_code["CE"]["response_count"] == 2
    assert by_code["CE"]["analytics_available"] is False
    ce_base = next(base for base in by_code["CE"]["bases"] if base["base_code"] == "FOR")
    assert ce_base["response_count"] == 2
    assert ce_base["analytics_available"] is False
    gp_base = next(base for base in by_code["GP"]["bases"] if base["base_code"] == "ANA")
    assert gp_base["response_count"] == 5


@pytest.mark.parametrize(
    "total,rows",
    [
        (6, [("GP", 5)]),  # one form has no Regional segment
        (8, [("GP", 5), ("UNLISTED", 3)]),  # non-catalog segment
    ],
)
def test_operational_total_remains_visible_when_a_segment_is_unassigned(monkeypatch, total, rows):
    configure_request(monkeypatch, total=total, regional_rows=rows)

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()

    assert payload["total_responses"] == total
    assert len(payload["regionals"]) == 10


def test_zero_forms_returns_zero_for_every_official_regional(monkeypatch):
    configure_request(monkeypatch, total=0)

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()

    assert payload["total_responses"] == 0
    assert payload["regionals_with_results"] == 0
    assert len(payload["regionals"]) == 10
    assert all(item["response_count"] == 0 for item in payload["regionals"])
    assert all(
        base["response_count"] == 0
        for regional in payload["regionals"]
        for base in regional["bases"]
    )


@pytest.mark.parametrize("count", [0, 1, 2, 3, 4, 5])
def test_operational_counts_show_exact_total_regional_and_base_values(monkeypatch, count):
    regional_rows = [("GP", count)] if count else []
    paired_rows = [("GP", "ANA", count)] if count else []
    configure_request(
        monkeypatch,
        total=count,
        regional_rows=regional_rows,
        paired_rows=paired_rows,
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    gp = next(row for row in payload["regionals"] if row["regional_code"] == "GP")
    bases = {base["base_code"]: base for base in gp["bases"]}

    assert payload["total_responses"] == count
    assert gp["response_count"] == count
    assert gp["analytics_available"] is (count >= 5)
    assert bases["ANA"]["response_count"] == count
    assert bases["ANA"]["analytics_available"] is (count >= 5)
    assert bases["BSB"]["response_count"] == 0


def test_operational_base_count_uses_its_stored_base_code_even_if_regional_pair_is_historical(monkeypatch):
    configure_request(
        monkeypatch,
        total=6,
        regional_rows=[("PR", 6)],
        paired_rows=[("PR", "BNU", 6)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    by_code = {row["regional_code"]: row for row in payload["regionals"]}
    bnu = next(base for base in by_code["SR"]["bases"] if base["base_code"] == "BNU")

    assert payload["total_responses"] == 6
    assert by_code["PR"]["response_count"] == 6
    assert by_code["SR"]["response_count"] == 0
    assert bnu["response_count"] == 6


def test_single_base_regional_displays_its_exact_operational_count(monkeypatch):
    configure_request(
        monkeypatch,
        total=2,
        regional_rows=[("PR", 2)],
        paired_rows=[("PR", "SJS", 2)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    pr = next(row for row in payload["regionals"] if row["regional_code"] == "PR")

    assert pr["response_count"] == 2
    assert len(pr["bases"]) == 1
    assert pr["bases"][0]["base_code"] == "SJS"
    assert pr["bases"][0]["response_count"] == 2


def test_small_and_complementary_base_counts_are_visible_as_operational_volume(monkeypatch):
    configure_request(
        monkeypatch,
        total=10,
        regional_rows=[("GP", 10)],
        paired_rows=[("GP", "ANA", 8), ("GP", "BSB", 2)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    gp = next(item for item in payload["regionals"] if item["regional_code"] == "GP")
    bases = {item["base_code"]: item for item in gp["bases"]}

    assert payload["total_responses"] == 10
    assert payload["regionals_with_results"] == 1
    assert gp["response_count"] == 10
    assert gp["analytics_available"] is True
    assert bases["ANA"]["response_count"] == 8
    assert bases["BSB"]["response_count"] == 2
    assert all("raw_count" not in item for item in gp["bases"])


def test_each_small_base_count_is_visible_for_authorized_operational_reporting(monkeypatch):
    configure_request(
        monkeypatch,
        total=5,
        regional_rows=[("GP", 5)],
        paired_rows=[("GP", "ANA", 2), ("GP", "BSB", 3)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    gp = next(item for item in payload["regionals"] if item["regional_code"] == "GP")
    bases = {item["base_code"]: item for item in gp["bases"]}

    assert payload["total_responses"] == 5
    assert gp["response_count"] == 5
    assert bases["ANA"]["response_count"] == 2
    assert bases["BSB"]["response_count"] == 3


def test_exact_operational_counts_are_not_suppressed_by_complementary_analytics_rules(monkeypatch):
    configure_request(
        monkeypatch,
        total=6,
        regional_rows=[("GP", 6)],
        paired_rows=[("GP", "ANA", 5), ("GP", "BSB", 1)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    gp = next(item for item in payload["regionals"] if item["regional_code"] == "GP")
    bases = {item["base_code"]: item for item in gp["bases"]}

    assert payload["total_responses"] == 6
    assert gp["response_count"] == 6
    assert bases["ANA"]["response_count"] == 5
    assert bases["BSB"]["response_count"] == 1


def test_operational_counts_remain_visible_for_stored_base_codes_with_historical_pairs(monkeypatch):
    configure_request(
        monkeypatch,
        total=6,
        regional_rows=[("PR", 6)],
        paired_rows=[("PR", "BNU", 6)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    by_code = {item["regional_code"]: item for item in payload["regionals"]}

    assert payload["total_responses"] == 6
    assert by_code["PR"]["response_count"] == 6
    assert by_code["SR"]["response_count"] == 0
    assert next(base for base in by_code["SR"]["bases"] if base["base_code"] == "BNU")["response_count"] == 6


def test_missing_base_segment_does_not_hide_global_or_regional_operational_counts(monkeypatch):
    configure_request(
        monkeypatch,
        total=6,
        regional_rows=[("GP", 6)],
        paired_rows=[("GP", None, 6)],
    )

    payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    gp = next(item for item in payload["regionals"] if item["regional_code"] == "GP")

    assert payload["total_responses"] == 6
    assert gp["response_count"] == 6
    assert all(base["response_count"] == 0 for base in gp["bases"])


@pytest.mark.parametrize("role", ["ADMIN", "MANAGEMENT", "SURVEY_ADMIN"])
def test_management_roles_can_export_adherence(monkeypatch, role):
    configure_request(monkeypatch, roles=(role,), total=0)

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence/export")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=\"adherence-CLIMATE_2026.xlsx\"" in response.headers["content-disposition"]


@pytest.mark.parametrize("access_type", ["INTERNAL", "EXTERNAL"])
def test_collaborator_cannot_export_adherence(monkeypatch, access_type):
    user = SimpleNamespace(id=17, status="ACTIVE", access_type=access_type)
    configure_request(monkeypatch, roles=("COLLABORATOR",))
    app.dependency_overrides[get_current_user] = lambda: user

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence/export")

    assert response.status_code == 403
    assert response.json() == {"detail": "MANAGEMENT_ACCESS_REQUIRED"}


@pytest.mark.parametrize("total", [0, 1, 2, 3, 4, 5])
def test_export_is_a_valid_localized_xlsx_with_exact_api_consistent_values(monkeypatch, total):
    paired_rows = []
    if total:
        first_base_count = total // 2
        second_base_count = total - first_base_count
        if first_base_count:
            paired_rows.append(("GP", "ANA", first_base_count))
        if second_base_count:
            paired_rows.append(("GP", "BSB", second_base_count))
    configure_request(
        monkeypatch,
        total=total,
        regional_rows=[("GP", total)] if total else [],
        paired_rows=paired_rows,
    )
    api_payload = client.get("/api/management/surveys/CLIMATE_2026/adherence").json()
    response = client.get(
        "/api/management/surveys/CLIMATE_2026/adherence/export?locale=zh"
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["区域", "服务中心"]
    regional_sheet, base_sheet = workbook["区域"], workbook["服务中心"]
    assert [cell.value for cell in regional_sheet[1]] == ["区域", "已完成问卷数", "已注册服务中心数"]
    assert [cell.value for cell in base_sheet[1]] == ["区域", "服务中心代码", "官方名称", "已完成问卷数"]
    assert regional_sheet.freeze_panes == "A2"
    assert base_sheet.freeze_panes == "A2"
    assert regional_sheet.auto_filter.ref == "A1:C11"
    assert base_sheet.auto_filter.ref == "A1:D33"
    assert regional_sheet[1][0].fill.fgColor.rgb.endswith("E60012")

    api_regionals = {item["regional_name"]: item for item in api_payload["regionals"]}
    exported_regionals = {
        regional_sheet.cell(row, 1).value: regional_sheet.cell(row, 2).value
        for row in range(2, 12)
    }
    for regional_name, regional in api_regionals.items():
        assert exported_regionals[regional_name] == regional["response_count"]

    exported_bases = {
        (base_sheet.cell(row, 1).value, base_sheet.cell(row, 2).value): base_sheet.cell(row, 4).value
        for row in range(2, 34)
    }
    for regional in api_payload["regionals"]:
        for base in regional["bases"]:
            key = (regional["regional_name"], base["base_code"])
            assert exported_bases[key] == base["response_count"]
    assert api_payload["total_responses"] == total
    assert exported_regionals["GP"] == total
    assert any(
        isinstance(cell.value, str) and "UTC" in cell.value
        for row in regional_sheet.iter_rows()
        for cell in row
    )
    assert not any(
        key in cell.value.lower()
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        for key in ("response_id", "user_id", "email", "answer_id")
    )


def test_export_zero_counts_remain_numeric_and_contains_all_official_bases(monkeypatch):
    configure_request(monkeypatch, total=0)

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence/export")
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    base_sheet = workbook["Base"]

    assert base_sheet.max_row >= 35
    assert {base_sheet.cell(row, 2).value for row in range(2, 34)} == {
        center.code for regional in REGIONAL_SC_CATALOG for center in regional.service_centers
    }
    assert all(base_sheet.cell(row, 4).value == 0 for row in range(2, 34))


def test_export_portuguese_workbook_uses_portuguese_headers_and_tabs(monkeypatch):
    configure_request(monkeypatch, total=0)

    response = client.get("/api/management/surveys/CLIMATE_2026/adherence/export?locale=pt")
    workbook = load_workbook(BytesIO(response.content), data_only=True)

    assert workbook.sheetnames == ["Regional", "Base"]
    assert [cell.value for cell in workbook["Regional"][1]] == [
        "Regional", "Formulários concluídos", "Bases cadastradas"
    ]
    assert [cell.value for cell in workbook["Base"][1]] == [
        "Regional", "Código da base", "Nome oficial da base", "Formulários concluídos"
    ]
